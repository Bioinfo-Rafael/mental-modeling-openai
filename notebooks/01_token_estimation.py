#!/usr/bin/env python3
"""Estimate existing 01 Notebook prompts locally; never call an inference API.

Run from any directory:
  .venv/bin/python notebooks/01_token_estimation.py --samples 10000 --seed 42
  .venv/bin/python notebooks/01_token_estimation.py --self-test

Code survey (2026-09-09): the actual directories are notebooks/ and tools/.
There is one notebooks/01_*.ipynb. It scans all official raw tasks, calls
preprocessing_for_inspection -> selected_query -> queries_for_sequence ->
build_prompt_queries, and uses count_query. Its batch selector is datasets='all'.
Task availability comes from the raw scan and llm_x.task.ALL_CLS; unsupported
task/version names remain unavailable, not aliased to another environment.
The raw tree is offline_data/<dataset>/raw_transitions/<task>/episodes/*.npz.
Each NPZ is an episode, array axis 0 is time. Adapters do not preprocess.
This script uses the SAME queries_for_sequence and count_query functions.
The user's explicit H override is the sole prompt-parameter change: Excel H
is actual contiguous history steps, so pass H-1 to the upstream H+1 builder.

Fixed/question counting removes exactly the history substring from the built
user message (retaining every surrounding separator), then tokenizes that
whole remaining user message. The system message is tokenized separately
because it is a separate API message, not an adjacent piece of user text.
The existing +9 framing estimate is assigned once to fixed/question.
Required additive totals are estimates: BPE boundary effects are exposed by
additional full-message counts. No output tokens, caching discounts, tools,
or guaranteed model-specific framing are included. mean+std is not a p95
or a confidence interval, nor the std of the complete prompt.

Only this script, its workbook, and the estimation dependency extra are
needed for delivery. Detailed sampling/prompt evidence is derived output
under outputs/token_estimation/ and is deliberately not committed.
"""
from __future__ import annotations

import argparse
import ast
from collections import defaultdict
from hashlib import sha256
from importlib.metadata import version
import io
import json
import math
from pathlib import Path
import random
import statistics
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.dont_write_bytecode = True

import numpy as np
import tiktoken
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from dataset_adapters.base import inventory_context
from dataset_adapters.registry import dataset_ids, get_adapter, EXTERNAL
from dataset_adapters.llmx import discover_episodes
from preprocessing.llmx_original import Episode, _history_range, _query_indices, file_sha256
from preprocessing.prompting import config_identity
from tools.count_input_tokens import count_query
from tools.inventory_llmx_data import create_inventory
from tools.live_inspection import official_frames
from tools.notebook_inspection import (
    preprocessing_for_inspection, queries_for_sequence, selected_query,
)
from tools.output_safety import derived_path

H_VALUES = [1, 10]  # Actual steps, NOT the upstream history_size parameter.
N_SAMPLES = 10000
RANDOM_SEED = 42
PRICE_CHECKED = "2026-09-09"
TOKENS_PER_MESSAGE = 3
PRIMING_TOKENS = 3
FRAMING = 2 * TOKENS_PER_MESSAGE + PRIMING_TOKENS
MODEL_DOCS = "https://developers.openai.com/api/docs/models/"

# Latest installed/PyPI tiktoken at investigation: 0.14.0.
# No fallback mapping is authorized without evidence. GPT-6 Astra is absent
# from that library and its model/token-counting docs do not specify encoding.
MODEL_ENCODING_MAP = {}  # Add only with a documented model-specific source.
MODELS = {
    "chatGPT3.5": dict(api_model="gpt-3.5-turbo", input_price_per_1m=0.50,
                       cached_input_price_per_1m=None, cache_write_price_per_1m=None, output_price_per_1m=1.50),
    "gpt-6-astra": dict(api_model="gpt-6-astra", input_price_per_1m=10.00,
                        cached_input_price_per_1m=1.00, cache_write_price_per_1m=12.50, output_price_per_1m=50.00),
    "gpt-5.6-sol": dict(api_model="gpt-5.6-sol", input_price_per_1m=4.00,
                        cached_input_price_per_1m=0.40, cache_write_price_per_1m=5.00, output_price_per_1m=20.00),
    "gpt-5.6-terra": dict(api_model="gpt-5.6-terra", input_price_per_1m=2.00,
                          cached_input_price_per_1m=0.20, cache_write_price_per_1m=2.50, output_price_per_1m=12.00),
    "gpt-5.6-luna": dict(api_model="gpt-5.6-luna", input_price_per_1m=0.20,
                         cached_input_price_per_1m=0.02, cache_write_price_per_1m=0.25, output_price_per_1m=1.20),
}
PRICE_SHEET = "Model_Prices"
PRICE_FIRST_ROW = 6
TASK_SHEET = "task description"
PRIORITY_TASKS = ['Acrobot-v1', 'MountainCar-v0', 'Pendulum-v1']


def task_order(task):
    return (PRIORITY_TASKS.index(task) if task in PRIORITY_TASKS else len(PRIORITY_TASKS), task)


HEADERS = [
    "タスク", "固定部分＋質問\n平均トークン数", "H=1 履歴\n平均トークン数",
    "H=1 履歴\n標準偏差（トークン）", "H=10 履歴\n平均トークン数", "H=10 合計\n平均トークン数",
    "H=10 合計金額\n平均（USD）", "H=10 履歴\n平均＋標準偏差（トークン）",
    "H=10 合計\n平均＋履歴標準偏差（トークン）", "H=10 合計金額\n平均＋履歴標準偏差（USD）",
    "H=10 履歴\n標準偏差（トークン）", "H=10 プロンプト全体\n平均トークン数",
    "H=10 分割集計－全体集計\n平均差（トークン）", "計測状況",
]
DESCRIPTION_LABELS = {
    'task': 'タスク', 'dataset': 'データセット', 'number_of_episodes': 'エピソード数',
    'total_timesteps': '総タイムステップ数', 'timesteps_mean': '平均エピソード長\n（ステップ）',
    'timesteps_min': '最小エピソード長\n（ステップ）', 'timesteps_max': '最大エピソード長\n（ステップ）',
    'state_dimensions': '状態の次元数', 'action_dimensions': '行動の次元数', 'reward_dimensions': '報酬の次元数',
    'state_shape': '状態配列の形状', 'action_shape': '行動配列の形状', 'reward_shape': '報酬配列の形状',
    'state_dtype': '状態のデータ型', 'action_dtype': '行動のデータ型', 'reward_dtype': '報酬のデータ型',
    'npz_keys': 'NPZ内の配列キー', 'raw_episode_files_size_bytes': 'エピソードファイル\n合計容量（バイト）',
    'schema_consistent_except_T': '時間軸を除く\nスキーマの一致', 'action_kind_observed': '行動の種別\n（観測値から推定）',
}
HEADER_ROW = 8
FIRST_ROW = HEADER_ROW + 1
NA = "n.a."
REPORT_MODELS = [name for name in MODELS if name != 'gpt-6-astra']
# Existing task_row statistics retain their internal schema. Only Excel display changes.
COLUMN_ORDER = [0, 6, 9, 1, 2, 3, 4, 5, 7, 8, 10, 11, 12, 13]
DISPLAY_HEADERS = [HEADERS[0], '100回分金額'] + [HEADERS[i] for i in COLUMN_ORDER[1:]]


def display_row(values):
    hundred = values[6] * 100 if isinstance(values[6], (int, float)) else NA
    return [values[0], hundred] + [values[i] for i in COLUMN_ORDER[1:]]


def resolve_models():
    """Use tiktoken's public mapping; fail visibly instead of generic fallback."""
    result = {}
    for label, settings in MODELS.items():
        api_model = settings["api_model"]
        try:
            encoding = tiktoken.encoding_for_model(api_model)
            resolution = "tiktoken.encoding_for_model"
            if api_model not in tiktoken.model.MODEL_TO_ENCODING:
                resolution += " (library prefix mapping, not model-specific API verification)"
        except KeyError:
            if api_model in MODEL_ENCODING_MAP:
                evidence = MODEL_ENCODING_MAP[api_model]
                if not evidence.get("source"):
                    raise ValueError("Explicit encoding requires a source")
                encoding = tiktoken.get_encoding(evidence["encoding"])
                resolution = evidence["source"]
            else:
                encoding = None
                resolution = "TOKENIZER_UNVERIFIED: no tiktoken mapping or official encoding evidence"
        result[label] = dict(**settings, encoding=encoding.name if encoding else None,
                             resolution=resolution, price_source=MODEL_DOCS + api_model,
                             price_checked=PRICE_CHECKED)
        print(f"Model: {label} (API ID: {api_model})\nTokenizer / encoding: "
              f"{encoding.name if encoding else NA}\nResolution: {resolution}\n"
              f"Standard uncached input USD/1M: {settings['input_price_per_1m']}", flush=True)
    return result


def notebook_scope():
    """Inspect, never execute or write, all actual first-party 01 notebooks."""
    paths = sorted(ROOT.glob("notebooks/01_*.ipynb"))
    if not paths:
        raise ValueError("No 01 Notebook found")
    facts = []
    for path in paths:
        nb = json.loads(path.read_text())
        text = "\n".join("".join(c["source"]) for c in nb["cells"] if c["cell_type"] == "code")
        tree = ast.parse(text)
        calls = [n for n in ast.walk(tree) if isinstance(n, ast.Call)]
        scopes = [ast.literal_eval(k.value) for c in calls
                  if isinstance(c.func, ast.Name) and c.func.id == "build_token_inventory"
                  for k in c.keywords if k.arg == "datasets"]
        if scopes != ["all"] or "selected_query(" not in text:
            raise ValueError(f"Notebook scope changed; review before estimating: {path}")
        # Do not silently ignore a future non-default metric/question/config.
        for n in ast.walk(tree):
            if isinstance(n, ast.Assign):
                for target in n.targets:
                    if isinstance(target, ast.Name) and target.id in {"METRIC", "QUESTION_NAME", "PREPROCESSING_CONFIG"}:
                        if ast.literal_eval(n.value) is not None:
                            raise ValueError(f"Review non-default {target.id} in {path}")
        facts.append(dict(path=str(path.relative_to(ROOT)), sha256=file_sha256(path),
                          code_sha256=sha256(text.encode()).hexdigest(), scope="all official tasks"))
    return facts


def choose_windows(sequences, h, n_samples, seed):
    """Uniform sample without replacement over all (file, valid target) pairs."""
    if h < 1 or n_samples < 1:
        raise ValueError("H and samples must be positive")
    pool = [(s.sequence_id, i) for s in sequences
            for i in _query_indices("next-action", h - 1, s.length)]
    selected = sorted(random.Random(seed).sample(pool, min(n_samples, len(pool))))
    assert len(selected) == len(set(selected))
    lengths = {s.sequence_id: s.length for s in sequences}
    for sid, index in selected:
        start, end = _history_range("next-action", index, h - 1)
        assert end - start == h and 0 <= start < end == index < lengths[sid]
    return len(pool), selected


def fixed_user_text(query):
    """Remove only the exact history span; retain all original separators."""
    if not query.history_text or query.user_prompt.count(query.history_text) != 1:
        raise ValueError("History must occur exactly once in the actual user prompt")
    prefix, history, suffix = query.user_prompt.partition(query.history_text)
    assert prefix + history + suffix == query.user_prompt
    assert query.question_text.strip() in suffix
    return prefix + suffix


def measure(query, encoding):
    counts = count_query(query, encoding, model="local-encoding:" + encoding.name,
                         encoding_name=encoding.name, fallback_used=False,
                         tokens_per_message=TOKENS_PER_MESSAGE, priming_tokens=PRIMING_TOKENS)
    fixed = counts["system_tokens"] + len(encoding.encode(fixed_user_text(query))) + FRAMING
    history = counts["history_tokens"]
    full = counts["estimated_api_input_tokens"]
    return dict(fixed=fixed, history=history, full=full, boundary_delta=fixed + history - full)


def stats(values):
    values = list(values)
    if not values:
        return dict(mean=None, std=None, min=None, max=None)
    return dict(mean=statistics.fmean(values), std=statistics.pstdev(values),
                min=min(values), max=max(values))


def estimate(h_values, n_samples, seed):
    notebooks = notebook_scope()
    models = resolve_models()
    encodings = {name: tiktoken.get_encoding(name) for name in {m["encoding"] for m in models.values()} if name}
    sources = discover_episodes()
    before = {str(s.path.relative_to(ROOT)): file_sha256(s.path) for s in sources}
    inventory = create_inventory(ROOT / "data/llmx_data")
    audit = dict(notebooks=notebooks, tiktoken_version=version("tiktoken"), models=models,
                 actual_h_values=h_values, upstream_h_values=[h-1 for h in h_values],
                 seed=seed, samples_per_task_h=n_samples, std_ddof=0, api_requests=0,
                 framing_estimate=FRAMING, raw_sha256=before, tasks={}, spot_checks=[])
    # Same live inventory / DataFrame entry point used by the 01 Notebook.
    audit['task_descriptions'] = task_descriptions(inventory)
    with inventory_context(inventory, {}):
        tasks = [name for name in dataset_ids() if name not in EXTERNAL]
        print("Detected tasks:", tasks)
        print("Detected data files:", len(sources), "\nNumber of episodes:", len(sources), flush=True)
        for task in tasks:
            adapter = get_adapter(task)
            sequences = adapter.sequences()
            config, reason = preprocessing_for_inspection(task)
            entry = dict(task=task, files=len(sequences), episodes=len(sequences),
                         status="ok" if config else "UNSUPPORTED_TASK: " + reason,
                         windows={}, encodings={name: {} for name in encodings})
            audit["tasks"][task] = entry
            print(f"\nTask: {task}; files/episodes: {len(sequences)}; {entry['status']}", flush=True)
            if config is None:
                continue
            for h in h_values:
                total, chosen = choose_windows(sequences, h, n_samples, seed)
                selection = defaultdict(set)
                for sid, i in chosen:
                    selection[sid].add(i)
                entry["windows"][h] = dict(valid=total, sampled=len(chosen),
                                           eligible_episodes=sum(bool(_query_indices("next-action", h-1, s.length)) for s in sequences),
                                           sampled_episodes=len(selection), selected=chosen)
                ctext, digest = config_identity(config, h-1)
                entry["windows"][h].update(preprocessing_config=ctext, preprocessing_config_sha256=digest)
                print(f"H={h} actual steps (upstream H={h-1}) valid windows: {total}; sampled windows: {len(chosen)}", flush=True)
                measures = {name: [] for name in encodings}
                spot_ids = {chosen[k] for k in [0, len(chosen)//2, len(chosen)-1]} if chosen else set()
                count = 0
                for seq in sequences:
                    needed = selection[seq.sequence_id]
                    if not needed:
                        continue
                    for q in queries_for_sequence(adapter, seq, config, h-1):
                        if q.query_index not in needed:
                            continue
                        assert q.metric == "next-action" and q.history_end-q.history_start == h
                        assert 0 <= q.history_start < q.history_end == q.query_index < seq.length
                        assert q.episode_path.resolve() == (ROOT / seq.source_file).resolve()
                        row_counts = {name: measure(q, encoding) for name, encoding in encodings.items()}
                        for name in measures:
                            measures[name].append(row_counts[name])
                        count += 1
                        if (seq.sequence_id, q.query_index) in spot_ids:
                            # Independent Notebook entry point, not just string-template tests.
                            reference = selected_query(adapter, seq, q.query_index, config, h-1)
                            assert q == reference
                            episode = Episode.load(q.episode_path)
                            assert q.history_text == episode.history_text(q.history_start, q.history_end,
                                indexed=True, drop_last_state_feature="Fetch" in task)
                            audit["spot_checks"].append(dict(task=task, actual_h=h, upstream_h=h-1,
                                source_file=seq.source_file, source_sha256=q.dataset_sha256,
                                episode_id=seq.episode_id, sequence_id=seq.sequence_id,
                                start=q.history_start, end_exclusive=q.history_end, target=q.query_index,
                                history=q.history_text, system_prompt=q.system_prompt, user_prompt=q.user_prompt,
                                fixed_user_text=fixed_user_text(q), question=q.question_text,
                                user_prompt_sha256=sha256(q.user_prompt.encode()).hexdigest(),
                                counts=row_counts, notebook_prompt_equal=True))
                        if q.query_index == max(needed):
                            break
                assert count == len(chosen)
                for name, rows in measures.items():
                    entry["encodings"][name][h] = {key: stats(r[key] for r in rows)
                                                          for key in ("history", "fixed", "full", "boundary_delta")}
                    values = entry["encodings"][name][h]
                    print(f"  {name}: fixed/question={values['fixed']['mean']}; history mean/std="
                          f"{values['history']['mean']}/{values['history']['std']}", flush=True)
    after = {str(s.path.relative_to(ROOT)): file_sha256(s.path) for s in discover_episodes()}
    assert before == after, "Raw data changed during estimation"
    assert all(file_sha256(ROOT / n["path"]) == n["sha256"] for n in notebooks), "Notebook changed during run"
    audit["raw_sha256_unchanged"] = True
    audit["notebooks_unchanged"] = True
    audit["code_revision"] = subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip()
    # Vendored source shares this repository's Git history. Record the original
    # upstream commit separately from local source content and project revision.
    audit["upstream_revision"] = json.loads((ROOT/'configs/sources.json').read_text())['llm_xavier']['commit']
    audit["upstream_source_sha256"] = {
        str(p.relative_to(ROOT/'upstream/LLM-Xavier')): file_sha256(p)
        for p in sorted((ROOT/'upstream/LLM-Xavier/llm_x').glob('*.py'))
    }
    return audit


def task_row(entry, model):
    status = entry["status"]
    if not model["encoding"]:
        status = (status + "; " if status != "ok" else "") + "TOKENIZER_UNVERIFIED"
    if status != "ok":
        return [entry["task"], *([NA]*12), status]
    values = entry["encodings"][model["encoding"]]
    one, ten = values[1], values[10]
    fixed, h1, s1, h10, s10 = ten['fixed']['mean'], one['history']['mean'], one['history']['std'], ten['history']['mean'], ten['history']['std']
    if any(v is None for v in [fixed, h1, s1, h10, s10]):
        return [entry["task"], *([NA]*12), "NO_VALID_WINDOW"]
    mean, high = fixed+h10, fixed+h10+s10
    price = model['input_price_per_1m']
    # Current standard rates are the <=272K tier; do not silently use this
    # sheet's simple requested cost formula for a future long-context run.
    if model['api_model'] != 'gpt-3.5-turbo' and max(mean, high, ten['full']['max']) > 272000:
        raise ValueError('Long-context price tier requires an explicit pricing review')
    return [entry['task'], fixed, h1, s1, h10, mean, mean/1e6*price,
            h10+s10, high, high/1e6*price, s10, ten['full']['mean'],
            round(mean-ten['full']['mean'], 9), 'ok']


def cache_formula_values(payload, values):
    """Store validated numeric formula caches in OOXML; formulas remain editable.

    openpyxl does not calculate formulas. Only the simple sums/products this
    script owns are cached, using independently computed task_row values.
    Excel recalculates on open. No formula text or raw input is overwritten.
    """
    namespace = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload)) as source, zipfile.ZipFile(result, 'w', zipfile.ZIP_DEFLATED) as dest:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename in values:
                tree = ET.fromstring(data)
                for c in tree.iter(namespace+'c'):
                    if c.find(namespace+'f') is not None:
                        value = values[info.filename][c.attrib['r']]
                        v = c.find(namespace+'v')
                        if v is None:
                            v = ET.SubElement(c, namespace+'v')
                        v.text = repr(float(value))
                data = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
            dest.writestr(info, data)
    return result.getvalue()


def task_descriptions(inventory):
    frames, _, _ = official_frames(inventory)
    frame_rows = {row['task']: row for row in frames.to_dict(orient='records')}

    def dimensions(values):
        return values[0] if len(values) == 1 and values[0] is not None else json.dumps(values) if values else NA

    def shapes(specs):
        return '; '.join(sorted({str(tuple(['T', *s['shape'][1:]])) if s['shape'] else '()' for s in specs})) or NA

    def dtypes(specs):
        return ', '.join(sorted({s['dtype'] for s in specs})) or NA

    result = []
    for task in inventory['tasks']:
        row = frame_rows[task['task']]
        reward_specs = task['rewards']['specs']
        result.append(dict(
            task=row['task'], dataset=task['dataset'],
            number_of_episodes=row['number_of_episodes'], total_timesteps=row['total_timesteps'],
            timesteps_mean=row['timesteps_mean'], timesteps_min=row['timesteps_min'], timesteps_max=row['timesteps_max'],
            state_dimensions=dimensions(row['state_dimensions']), action_dimensions=dimensions(row['action_dimensions']),
            reward_dimensions=dimensions(sorted({math.prod(s['shape'][1:]) for s in reward_specs if s['shape']})),
            state_shape=shapes(task['states']['specs']), action_shape=shapes(task['actions']['specs']), reward_shape=shapes(reward_specs),
            state_dtype=dtypes(task['states']['specs']), action_dtype=dtypes(task['actions']['specs']), reward_dtype=dtypes(reward_specs),
            npz_keys=', '.join(row['npz_keys']), raw_episode_files_size_bytes=task['raw_episode_files_size_bytes'],
            schema_consistent_except_T=row['schema_consistent_except_T'], action_kind_observed=', '.join(row['action_kind'])))
    return sorted(result, key=lambda row: task_order(row['task']))


def write_task_descriptions(wb, descriptions):
    descriptions = sorted(descriptions, key=lambda row: task_order(row['task']))
    ws = wb.create_sheet(TASK_SHEET)
    ws.sheet_view.showGridLines = False
    ws['A2'] = 'Task description: official raw episodes'
    ws['A3'] = 'All available episodes, not sampled windows. Raw dimensions are before preprocessing.'
    summary = ['task', 'dataset', 'number_of_episodes', 'total_timesteps', 'timesteps_mean', 'timesteps_min',
               'timesteps_max', 'state_dimensions', 'action_dimensions', 'reward_dimensions']
    detail = ['task', 'state_shape', 'action_shape', 'reward_shape', 'state_dtype', 'action_dtype', 'reward_dtype',
              'npz_keys', 'raw_episode_files_size_bytes', 'schema_consistent_except_T', 'action_kind_observed']
    detail_start = 9 + len(descriptions)
    for header_row, fields in [(5, summary), (detail_start, detail)]:
        for c, name in enumerate(fields, 1):
            ws.cell(header_row, c, DESCRIPTION_LABELS[name])
        ws.row_dimensions[header_row].height = 44
        for r, entry in enumerate(descriptions, header_row+1):
            ws.row_dimensions[r].height = 42 if header_row == detail_start else 25
            for c, name in enumerate(fields, 1):
                cell = ws.cell(r, c, entry[name])
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                if isinstance(entry[name], (int, float)) and not isinstance(entry[name], bool):
                    cell.number_format = '#,##0.00' if name == 'timesteps_mean' else '#,##0'
    notes = [
        'Source: tools.inventory_llmx_data.create_inventory -> tools.live_inspection.official_frames, as in notebooks/01_data_and_token_inspection.ipynb.',
        'Raw files: data/llmx_data/offline_data/<dataset>/raw_transitions/<task>/episodes/*.npz. One file = one episode.',
        'Episode length = states.shape[0]. Total = sum of lengths. Mean = arithmetic mean across episodes; min/max use the same population.',
        'Dimensions = product of shape[1:] (all non-time axes), not ndarray.ndim. Discrete action dimension 1 is not the number of possible actions.',
        'T denotes the episode time axis. Singleton axes are retained. All observed shape/dtype variants are listed; multiple dimensions appear as a list.',
        'Raw dimensions differ from processed input selection (e.g. Fetch feature selection). No raw arrays are modified.',
        'action_kind_observed is the inventory heuristic: a scalar integer-valued action is labelled discrete. This is not verified environment action-space metadata.',
        'raw_episode_files_size_bytes is NPZ container size. schema_consistent_except_T checks every NPZ key/shape/dtype with the time length abstracted as T.',
    ]
    for r, note in enumerate(notes, detail_start+len(descriptions)+3):
        ws.cell(r, 1, note)
    for row in ws:
        for cell in row:
            cell.font = Font(name='Arial', size=10)
    ws['A2'].font = Font(name='Arial', size=14, bold=True)
    for row_number, count in [(5, len(summary)), (detail_start, len(detail))]:
        for cell in ws[row_number][:count]:
            cell.fill = PatternFill('solid', fgColor='283C58')
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for c, width in enumerate([34, 23, 23, 22, 20, 20, 20, 35, 29, 29, 23], 1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = 'B6'
    ws.auto_filter.ref = f'A5:J{5+len(descriptions)}'


def write_prices(wb, models):
    """One authoritative editable price table; result sheets link to column C."""
    ws = wb.create_sheet(PRICE_SHEET)
    ws.sheet_view.showGridLines = False
    ws['A2'] = 'Model prices (USD / 1M tokens)'
    ws['A3'] = 'Edit yellow price cells. Cost columns use Input only; token statistics do not change.'
    headers = ['モデル', 'APIモデルID', '入力単価', 'キャッシュ入力単価', 'キャッシュ書込単価', '出力単価', '確認日', '', '公式情報の参照先']
    for c, value in enumerate(headers, 1):
        ws.cell(5, c, value)
    for r, (label, model) in enumerate(models.items(), PRICE_FIRST_ROW):
        fields = ['input_price_per_1m', 'cached_input_price_per_1m', 'cache_write_price_per_1m', 'output_price_per_1m']
        values = [label, model['api_model'], *[model[k] if model[k] is not None else NA for k in fields],
                  model['price_checked'], '', model['price_source']]
        for c, value in enumerate(values, 1):
            ws.cell(r, c, value)
        ws.row_dimensions[r].height = 26
        for c in range(3, 7):
            cell = ws.cell(r, c)
            cell.number_format = '"$"0.00'
            cell.alignment = Alignment(horizontal='right', vertical='center')
            if isinstance(cell.value, (int, float)):
                cell.fill = PatternFill('solid', fgColor='FFF0CD')
    ws['A12'] = 'Standard rates. GPT-5.6/GPT-6 rows apply to requests up to 272K input tokens. No Batch/Flex/Fast adjustments.'
    ws['A13'] = 'Cache write = published 1.25x Input rate for GPT-5.6/GPT-6. GPT-3.5 cache rates are not listed (n.a.), not zero.'
    ws['A14'] = 'Astra price is retained for reference; its estimate sheet is omitted at user request.'
    ws['A15'] = 'Keep model rows in place; formulas use direct row references. Rerunning the script restores the prices in MODELS.'
    for row in ws:
        for cell in row:
            cell.font = Font(name='Arial', size=10)
    ws['A2'].font = Font(name='Arial', size=14, bold=True)
    ws.row_dimensions[5].height = 30
    for c in [1, 2, 3, 4, 5, 6, 7, 9]:
        cell = ws.cell(5, c)
        cell.fill = PatternFill('solid', fgColor='283C58')
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for c, width in enumerate([20, 22, 14, 16, 16, 14, 16, 3, 77], 1):
        ws.column_dimensions[get_column_letter(c)].width = width


def write_workbook(audit, output):
    wb = Workbook()
    wb.remove(wb.active)
    expected = {}
    cached = {}
    tasks = sorted(audit['tasks'], key=task_order)
    for sheet_no, label in enumerate(REPORT_MODELS, 1):
        model = audit['models'][label]
        ws = wb.create_sheet(label)
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 85
        ws['A2'] = label
        ws['A2'].font = Font(name='Arial', size=14, bold=True)
        price_ref = f"'{PRICE_SHEET}'!$C${PRICE_FIRST_ROW + list(audit['models']).index(label)}"
        ws['A3'], ws['B3'] = 'Input USD / 1M', f'={price_ref}'
        ws['D3'], ws['E3'] = 'API model', model['api_model']
        ws['A4'], ws['B4'] = 'Encoding', model['encoding'] or NA
        ws['D4'], ws['E4'] = 'tiktoken', audit['tiktoken_version']
        ws['G3'], ws['H3'] = 'Price checked', PRICE_CHECKED
        ws['G4'], ws['H4'] = 'Seed / max samples', f"{audit['seed']} / {audit['samples_per_task_h']}"
        ws['A6'] = 'H = actual history steps. Upstream H = H-1. Input only; standard uncached rates; +9 framing is an estimate.'
        for col, header in enumerate(DISPLAY_HEADERS, 1):
            ws.cell(HEADER_ROW, col, header)
        rows = [task_row(audit['tasks'][task], model) for task in tasks]
        expected[label] = rows
        cache = {'B3': model['input_price_per_1m']}
        for r, values in enumerate(rows, FIRST_ROW):
            displayed = display_row(values)
            for c, value in enumerate(displayed, 1):
                ws.cell(r, c, value)
            if values[-1] == 'ok':
                formulas = {2:f'=C{r}*100', 3:f'=I{r}/1000000*{price_ref}',
                            4:f'=K{r}/1000000*{price_ref}', 9:f'=E{r}+H{r}',
                            10:f'=H{r}+L{r}', 11:f'=E{r}+J{r}', 14:f'=ROUND(I{r}-M{r},9)'}
                for c, formula in formulas.items():
                    ws.cell(r, c, formula)
                    cache[f'{get_column_letter(c)}{r}'] = displayed[c-1]
                print(f"{label} / {values[0]}: H10 total mean={values[5]:.2f}; mean+std={values[8]:.2f}; "
                      f"estimated cost=${values[6]:.8f} / ${values[9]:.8f}", flush=True)
        cached[f'xl/worksheets/sheet{sheet_no}.xml'] = cache
        last = FIRST_ROW + len(tasks)-1
        ws.freeze_panes = 'E9'
        ws.auto_filter.ref = f'A8:O{last}'
        ws.print_title_rows = '1:8'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        widths = [34, 23, 24, 30, 23, 21, 23, 23, 24, 30, 32, 23, 27, 32, 46]
        for c, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.row_dimensions[HEADER_ROW].height = 74
        for r in range(FIRST_ROW, last+1):
            ws.row_dimensions[r].height = 29
            for c in range(2, 15):
                cell = ws.cell(r,c)
                cell.number_format = '"$"0.00000000' if c in (2,3,4) else '#,##0.00'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            ws.cell(r,15).alignment = Alignment(wrap_text=True, vertical='center')
            if rows[r-FIRST_ROW][-1] != 'ok':
                ws.cell(r,15).fill = PatternFill('solid', fgColor='FFF0CD')
                ws.row_dimensions[r].height = 58
        # Reproducibility/coverage is below the requested 10-column output.
        base = last + 4
        scope_headers = ['タスク', '実際の履歴長H\n（ステップ）', '公式関数へ渡すH', '有効ウィンドウ数', '抽出ウィンドウ数',
                         '抽出対象エピソード数', '抽出済みエピソード数', 'rawファイル数', '標準偏差のddof', '固定部分＋質問\n標準偏差（トークン）']
        for c,v in enumerate(scope_headers,1): ws.cell(base,c,v)
        line = base+1
        for task in tasks:
            entry = audit['tasks'][task]
            for h in audit['actual_h_values']:
                scope = entry['windows'].get(h,{})
                fixed_stats = entry['encodings'].get(model['encoding'],{}).get(h,{}).get('fixed',{})
                vals = [task,h,h-1,scope.get('valid',NA),scope.get('sampled',NA),
                        scope.get('eligible_episodes',NA),scope.get('sampled_episodes',NA),
                        entry['files'],0,fixed_stats.get('std',NA)]
                for c,v in enumerate(vals,1): ws.cell(line,c,v)
                line += 1
        notes = [
            ('Method', 'Uniform random sample without replacement over ALL valid (episode, target index) pairs per task/H. Windows never cross files.'),
            ('History', 'Excel H=1/10 means 1/10 consecutive steps, using existing Notebook builder with history_size=0/9. Labels and numeric precision unchanged.'),
            ('Fixed/question', 'Tokenize whole user message with history span removed; preserve prefix, suffix, question and all separators. System stays a separate message. Add +9 once.'),
            ('Additive total', 'Required total = fixed/question mean + history mean. BPE boundaries can differ from counting complete user text; column M counts full messages, N shows the difference.'),
            ('100回分金額', 'B列 = C列の平均入力金額（USD） x 100。同じ条件の入力100回分。キャッシュ割引・出力料金は含まない。'),
            ('Mean + std', 'Population std (ddof=0) of complete history windows. Not a confidence interval, not p95, and not std of total prompt.'),
            ('Pricing', 'One input request; standard uncached input, no output/tools/cache fees. The +9 framing is inherited from count_query, not model-specific API usage.'),
            ('Unavailable', 'n.a. is deliberate: unsupported upstream task or unverified tokenizer. No zero filling, proxy task, or guessed encoding.'),
            ('Tokenizer evidence', model['resolution']),
            ('Tokenizer code', 'tiktoken 0.14.0 model.py: GPT-3.5 exact mapping; GPT-5.6 uses library gpt-5 prefix mapping. GPT-6 mapping absent.'),
            ('Price source', 'Model_Prices sheet: editable rates, check date and official sources. Cost = total tokens / 1000000 * Input price.'),
            ('Counting limits', 'https://developers.openai.com/api/docs/guides/token-counting'),
            ('Notebook', ', '.join(n['path'] for n in audit['notebooks'])),
            ('Reused functions', 'queries_for_sequence / selected_query / count_query; upstream Episode.history_text / render_question / system_prompt / user_prompt.'),
            ('Data', 'data/llmx_data/offline_data/<dataset>/raw_transitions/<task>/episodes/*.npz; one file = one episode.'),
            ('Reproduce', '.venv/bin/python notebooks/01_token_estimation.py --samples 10000 --seed 42 --h-values 1 10'),
            ('Evidence', 'outputs/token_estimation/audit.json includes source hashes, selected windows, config hashes and representative exact prompt strings.'),
        ]
        for title, value in notes:
            line += 1
            ws.cell(line,1,title)
            ws.cell(line,2,value)
        for row in ws:
            for cell in row:
                if cell.coordinate != 'A2': cell.font = Font(name='Arial',size=10)
                if cell.alignment.vertical is None: cell.alignment = Alignment(vertical='center')
        for header_row in [HEADER_ROW,base]:
            ws.row_dimensions[header_row].height = 74 if header_row == HEADER_ROW else 46
            for cell in ws[header_row][:15 if header_row==HEADER_ROW else 10]:
                cell.fill = PatternFill('solid',fgColor='283C58')
                cell.font = Font(name='Arial',size=10,bold=True,color='FFFFFF')
                cell.alignment = Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws['B3'].number_format = '"$"0.00'
        for col in (2, 3, 4):
            ws.cell(HEADER_ROW, col).fill = PatternFill('solid', fgColor='9C5700')
    write_prices(wb, audit['models'])
    write_task_descriptions(wb, audit['task_descriptions'])
    payload = io.BytesIO()
    wb.save(payload)
    output = derived_path(output)
    output.parent.mkdir(parents=True,exist_ok=True)
    output.write_bytes(cache_formula_values(payload.getvalue(),cached))
    validate_workbook(output, expected)
    return expected


def validate_workbook(path, expected):
    wb = load_workbook(path, data_only=True)
    formula_wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == [*REPORT_MODELS, PRICE_SHEET, TASK_SHEET]
    task_lists = []
    for label, rows in expected.items():
        price_row = PRICE_FIRST_ROW + list(MODELS).index(label)
        ws = wb[label]
        price_ref = f"'{PRICE_SHEET}'!$C${price_row}"
        assert wb[PRICE_SHEET].cell(price_row, 1).value == label
        assert wb[PRICE_SHEET].cell(price_row, 3).value == ws['B3'].value
        assert formula_wb[label]['B3'].value == f'={price_ref}'
        assert [c.value for c in ws[HEADER_ROW][:len(DISPLAY_HEADERS)]] == DISPLAY_HEADERS
        assert ws.freeze_panes == 'E9' and ws.auto_filter.ref
        tasks = []
        for r, expected_row in enumerate(rows,FIRST_ROW):
            displayed = [ws.cell(r,c).value for c in range(1,len(DISPLAY_HEADERS)+1)]
            actual = [displayed[0]] + [displayed[1+COLUMN_ORDER.index(i)] for i in range(1,len(HEADERS))]
            tasks.append(actual[0])
            for x,y in zip(actual,expected_row):
                if isinstance(y,(int,float)):
                    assert math.isfinite(x) and math.isclose(x,y,rel_tol=1e-12,abs_tol=1e-9),(label,r,x,y)
                else: assert x==y,(label,r,x,y)
            if actual[-1]=='ok':
                _,f,_,_,h,total,cost,high,total_high,cost_high,std,_,_,_ = actual
                assert high >= h and math.isclose(high,h+std)
                assert math.isclose(total,f+h) and math.isclose(total_high,f+high)
                price = ws['B3'].value
                assert math.isclose(cost,total/1e6*price) and math.isclose(cost_high,total_high/1e6*price)
                assert math.isclose(displayed[1], cost*100)
                assert formula_wb[label].cell(r,2).value == f'=C{r}*100'
                assert formula_wb[label].cell(r,3).value == f'=I{r}/1000000*{price_ref}'
                assert formula_wb[label].cell(r,4).value == f'=K{r}/1000000*{price_ref}'
        task_lists.append(tasks)
    assert all(t==task_lists[0] for t in task_lists)
    assert task_lists[0][:3] == PRIORITY_TASKS
    for label in REPORT_MODELS:
        assert all(formula_wb[label].cell(8,c).fill.fgColor.rgb == '009C5700' for c in (2,3,4))
    assert [wb[TASK_SHEET].cell(r, 1).value for r in range(6, 6+len(task_lists[0]))] == task_lists[0]
    print('Validation: four model sheets plus Model_Prices and task description; 100-query costs, all tasks, cross-sheet formulas/caches and panes verified.',flush=True)


def self_test():
    from dataset_adapters.base import Sequence
    seqs = [Sequence(0,'a','a',30,'.npz'),Sequence(1,'b','b',40,'.npz')]
    for h in [1,10,15]:
        n, chosen = choose_windows(seqs,h,100000,42)
        assert n == sum(s.length-h for s in seqs)
        assert {sid for sid,_ in chosen} == {0,1}
        assert choose_windows(seqs,h,7,42)==choose_windows(seqs,h,7,42)
    try: choose_windows(seqs,0,10,42)
    except ValueError: pass
    else: raise AssertionError('H=0 must be rejected')
    inv = create_inventory(ROOT/'data/llmx_data')
    descriptions = {row['task']: row for row in task_descriptions(inv)}
    for task in inv['tasks']:
        row = descriptions[task['task']]
        lengths = [f['arrays']['states']['shape'][0] for f in task['files']]
        assert row['number_of_episodes'] == len(task['files'])
        assert row['total_timesteps'] == sum(lengths)
        assert row['timesteps_mean'] == statistics.fmean(lengths)
        assert row['timesteps_min'] == min(lengths) and row['timesteps_max'] == max(lengths)
        for field, key in [('state_dimensions', 'states'), ('action_dimensions', 'actions')]:
            dims = sorted({math.prod(f['arrays'][key]['shape'][1:]) for f in task['files']})
            assert row[field] == (dims[0] if len(dims) == 1 else json.dumps(dims))
    with inventory_context(inv,{}):
        adapter = get_adapter('MountainCar-v0')
        config,_ = preprocessing_for_inspection(adapter.dataset_id)
        seq=adapter.sequences()[0]
        for h in [1,10]:
            q = selected_query(adapter,seq,20,config,h-1)
            assert q.history_start==20-h and q.history_end==20
            assert q.history_text.count('Step ')==h
            for name in ['cl100k_base','o200k_base']:
                row=measure(q,tiktoken.get_encoding(name))
                assert row['fixed']+row['history']-row['full']==row['boundary_delta']
    assert stats([1,3])['std']==1
    print('Self-tests passed: exact H, boundaries, two files, reproducible sampling, existing prompt builder, token decomposition.')


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--samples',type=int,default=N_SAMPLES)
    parser.add_argument('--seed',type=int,default=RANDOM_SEED)
    parser.add_argument('--h-values',type=int,nargs='+',default=H_VALUES)
    parser.add_argument('--output',type=Path,default=ROOT/'notebooks/01data/token_estimation.xlsx')
    parser.add_argument('--audit-dir',type=Path,default=ROOT/'outputs/token_estimation')
    parser.add_argument('--self-test',action='store_true')
    args=parser.parse_args()
    if args.self_test:
        self_test()
        return
    hs=sorted(set(args.h_values))
    if args.samples<1 or not hs or min(hs)<1 or not {1,10}<=set(hs):
        parser.error('Positive sample count and H values required; include H=1 and H=10 for the requested columns.')
    output=derived_path(args.output)
    audit_dir=derived_path(args.audit_dir)
    result=estimate(hs,args.samples,args.seed)
    expected=write_workbook(result,output)
    result['rows']=expected
    result['validation']='passed; deliberate n.a. retained for unsupported tasks/unverified tokenizer'
    audit_dir.mkdir(parents=True,exist_ok=True)
    (audit_dir/'audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
    print('Spot checks:',len(result['spot_checks']),'exact Notebook prompt matches; raw SHA256 unchanged.')
    for h in [1,10]:
        sample=next(s for s in result['spot_checks'] if s['task']=='MountainCar-v0' and s['actual_h']==h)
        print('Spot check:',sample['source_file'],'episode',sample['episode_id'],
              'history',sample['start'],sample['end_exclusive'],'target',sample['target'])
        print(sample['history'])
        print('Tokens:',sample['counts'])
    print('Saved:',output,'\nAudit:',audit_dir/'audit.json')


if __name__=='__main__':
    main()
